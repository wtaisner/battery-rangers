"""A script to generate an Excel file with molecular structures and SMILES strings."""
import argparse
import os
import shutil
from ast import literal_eval

import pandas as pd
from openpyxl import Workbook
from openpyxl.drawing.image import Image
from rdkit import Chem
from rdkit.Chem import Draw


def read_model_data(input_file: str) -> dict:
    """
    Function to read SMILES data from a file.
    The function should parse the file and return a dictionary
    in the format: {"Model_Name": ["SMILES1", "SMILES2", ...], ...}

    Args:
        input_file (str): Path to the input file.

    Returns:
        dict: Dictionary of models and their SMILES strings.
    """
    data = pd.read_csv(input_file)
    data["smiles_after_filtering"] = data["smiles_after_filtering"].apply(literal_eval)
    # change df to dict model_name: smiles_after_filtering
    data = data.set_index("model_name")["smiles_after_filtering"].to_dict()
    return data


def generate_excel(model_data: dict, output_file: str, output_folder: str):
    """
    Generates an Excel file with SMILES structures and images.

    Args:
        model_data (dict): Dictionary of models and their SMILES strings.
        output_file (str): Path to save the output Excel file.
        output_folder (str): Path to the folder for temporary images.

    Returns:
        None
    """
    # Create a new Excel workbook
    wb = Workbook()

    # Create the output folder for temporary images
    os.makedirs(output_folder, exist_ok=True)

    for model_name, smiles_list in model_data.items():
        # Create a new sheet for each model
        ws = wb.create_sheet(title=model_name)

        # Add headers
        ws.cell(row=1, column=1, value="SMILES")
        ws.cell(row=1, column=2, value="Structure")
        ws.cell(row=1, column=3, value="Grade (0-5)")

        # Set default column widths
        ws.column_dimensions["A"].width = 30  # SMILES column
        ws.column_dimensions["B"].width = 50  # Image column
        ws.column_dimensions["C"].width = 20  # Grade column

        # Set default row height
        for row in range(2, len(smiles_list) + 2):  # Adjust for data rows
            ws.row_dimensions[row].height = 250

        # Process each SMILES in the list
        for i, smiles in enumerate(smiles_list, start=2):
            # Parse the molecule
            mol = Chem.MolFromSmiles(smiles)

            # Add SMILES to the Excel
            ws.cell(row=i, column=1, value=smiles)

            img_path = os.path.join(output_folder, f"{model_name}_mol_{i}.png")
            Draw.MolToFile(mol, img_path)

            # Add the image to the Excel
            img = Image(img_path)
            ws.add_image(img, f"B{i}")

            # Add grading instructions
            ws.cell(row=i, column=3, value="Enter 0-5")

    # Remove the default sheet
    if "Sheet" in wb.sheetnames:
        del wb["Sheet"]

    # Save the Excel file
    wb.save(output_file)
    print(f"Excel file saved as {output_file}.")


def main():
    """
    Main function to parse arguments and generate the Excel file.
    """
    parser = argparse.ArgumentParser(description="Generate an Excel file with molecular structures and SMILES strings.")
    parser.add_argument("-i", "--input_file", help="Path to the input file containing model data.")
    parser.add_argument("-o", "--output_file", default="Molecules_with_Models.xlsx", help="Path to save the output Excel file.")
    parser.add_argument("-t", "--temp_folder", default="temp_images", help="Folder to store temporary images.")
    args = parser.parse_args()

    # Read the data from the input file
    model_data = read_model_data(args.input_file)
    if not model_data:
        print("Error: No data loaded from the input file.")
        return

    # Generate the Excel file
    try:
        generate_excel(model_data, args.output_file, args.temp_folder)
    finally:
        # Clean up temporary folder
        if os.path.exists(args.temp_folder):
            shutil.rmtree(args.temp_folder)
            print(f"Temporary folder '{args.temp_folder}' has been deleted.")


if __name__ == "__main__":
    main()
