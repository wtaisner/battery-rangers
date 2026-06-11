import argparse
import io
import math
import os

import pandas as pd
from openpyxl import load_workbook
from openpyxl.drawing.image import Image as XLImage
from rdkit import Chem
from rdkit.Chem import Draw


def generate_filtered_excel(input_file, output_folder="filtered_batches_rdkit", chunk_size=250):
    print(f"Reading {input_file}...")

    # 1. Read and Filter Data
    # -----------------------
    try:
        df = pd.read_excel(input_file)
    except Exception as e:
        print(f"Error reading file: {e}")
        return

    smiles_col = df.columns[0]
    # Identify filter columns (last 3)
    if df.shape[1] > 2:
        print("Filtering columns present")

        # Assume Column 0 is SMILES, Column 1 is 'molecule' (placeholder), last 3 are filters
        filter_cols = df.columns[-3:]

        print(f"Using '{smiles_col}' as SMILES column.")
        print(f"Filtering on: {list(filter_cols)}")

        # Apply 'no' filter (case-insensitive)
        mask = df[filter_cols].astype(str).apply(lambda x: x.str.strip().str.lower()) == "no"
        df = df[mask.all(axis=1)].copy()

    total_rows = len(df)
    print(f"Rows passing filter: {total_rows}")

    if total_rows == 0:
        return

    # Ensure the output directory exists
    os.makedirs(output_folder, exist_ok=True)

    # 2. Process in Batches
    # ---------------------
    num_batches = math.ceil(total_rows / chunk_size)

    for i in range(num_batches):
        start_idx = i * chunk_size
        end_idx = start_idx + chunk_size

        # Get the batch dataframe
        batch_df = df.iloc[start_idx:end_idx].copy()

        # Define output filename
        filename = os.path.join(output_folder, f"batch_{i + 1}.xlsx")

        # Step A: Save text data first
        # We save index=False so Col A is SMILES, Col B is molecule, etc.
        batch_df.to_excel(filename, index=False)

        # Step B: Open file with OpenPyXL to insert RDKit images
        wb = load_workbook(filename)
        ws = wb.active

        # Adjust Column B width to fit images (approx 40 units ~ 300px)
        ws.column_dimensions["B"].width = 38

        print(f"Processing batch {i + 1}/{num_batches} (generating images)...")

        # Iterate through rows in the batch
        # Enumerate gives us 0-based index, but Excel rows start at 2 (1 is header)
        for row_idx, smiles in enumerate(batch_df[smiles_col]):
            excel_row_num = row_idx + 2

            # Generate Image with RDKit
            mol = Chem.MolFromSmiles(str(smiles))

            if mol:
                try:
                    # Create image data in memory
                    img = Draw.MolToImage(mol, size=(300, 300))

                    # Convert PIL image to BytesIO for OpenPyXL
                    img_byte_arr = io.BytesIO()
                    img.save(img_byte_arr, format="PNG")
                    img_byte_arr.seek(0)

                    # Create OpenPyXL Image
                    xl_img = XLImage(img_byte_arr)

                    # Position image in Column B
                    anchor = f"B{excel_row_num}"
                    ws.add_image(xl_img, anchor)

                    # Set row height to accommodate image (approx 225 points ~ 300px)
                    ws.row_dimensions[excel_row_num].height = 225

                    # Optional: Clear text from the cell behind the image
                    ws[anchor] = ""

                except Exception as e:
                    print(f"Failed to draw SMILES at row {excel_row_num}: {e}")
            else:
                # If invalid SMILES, write text instead of image
                ws[f"B{excel_row_num}"] = "Invalid SMILES"

        # Save the workbook with images
        wb.save(filename)
        print(f"Saved {filename}")

    print("\nProcessing complete.")


if __name__ == "__main__":
    parser = argparse.ArgumentParser()
    parser.add_argument("--input", required=True, help="Path to input Excel file")
    args = parser.parse_args()

    generate_filtered_excel(args.input)
