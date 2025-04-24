import marimo

__generated_with = "0.13.1"
app = marimo.App()


@app.cell
def _():
    import re
    from ast import literal_eval

    import marimo as mo
    import pandas as pd

    from modules.core.molecule_filter import MoleculeFilter

    filter = MoleculeFilter()
    return literal_eval, mo, pd, re


@app.cell(hide_code=True)
def _(mo):
    mo.md(r"""# Format results of sampling & filtering""")
    return


@app.cell
def _(literal_eval, pd):
    df = pd.read_csv("data/sampling/old/filtering_comparison.csv")
    df["smiles_after_filtering"] = df["smiles_after_filtering"].apply(literal_eval)
    df.shape
    return (df,)


@app.cell
def _(df):
    df.head()
    return


@app.cell
def _(Chem, pd):
    experts1 = pd.read_csv("data/raw/data_experts_1.csv")
    experts2 = pd.read_csv("data/raw/data_experts2.csv")
    experts3 = pd.read_csv("data/raw/data_experts3.csv")

    # canonicalize smiles for all experts
    experts1["smiles"] = experts1["smiles"].apply(lambda x: Chem.CanonSmiles(x))
    experts2["smiles"] = experts2["smiles"].apply(lambda x: Chem.CanonSmiles(x))
    experts3["smiles"] = experts3["smiles"].apply(lambda x: Chem.CanonSmiles(x))
    return experts1, experts2, experts3


@app.cell
def _(experts1, experts2, experts3, pd):
    df_experts = pd.concat([experts1, experts2, experts3], ignore_index=True)["smiles"]
    print(df_experts.shape)
    df_experts = set(df_experts.to_list())
    len(df_experts)
    return (df_experts,)


@app.cell
def _(df, pd, re):
    def parse_filenames(df, column_name):
        """
        Parse a column of filenames into separate features.

        Args:
            df (pd.DataFrame): Input DataFrame.
            column_name (str): Name of the column to parse.

        Returns:
            pd.DataFrame: DataFrame with new feature columns.
        """

        def parse_row(row):
            pattern = {
                "model_name": "(bionemo_[a-z]+|reinvent|filtered|symmetrical)",
                "seed": "(data_experts_\\d+)",
                "num_samples": "num_samples_(\\d+)|(\\d+smiles)",
                "sampling_method": "sampling_method_([\\w-]+)",
                "scaled_radius": "scaled_radius_(\\d+(\\.\\d+)?)",
                "chunks": "(\\d+x\\d+)_chunks",
                "beam_size": "beam_size_(\\d+)",
                "beam_alpha": "beam_alpha_(\\d+(\\.\\d+)?)",
                "top_k": "top_k_(\\d+)",
                "top_p": "top_p_(\\d+(\\.\\d+)?)",
                "temperature": "temperature_(\\d+(\\.\\d+)?)",
            }
            parsed_data = {}
            for key, regex in pattern.items():
                match = re.search(regex, row)
                if key == "model_name":
                    parsed_data[key] = match.group(1).replace("bionemo_", "") if match and "bionemo_" in match.group(1) else match.group(1)
                    if parsed_data[key] == "filtered" or parsed_data[key] == "symmetrical":
                        parsed_data[key] = "reinvent"
                elif key == "seed":
                    parsed_data[key] = match.group(1) if match else None
                elif key == "num_samples":
                    if match:
                        num_match = match.group(1) or match.group(2)
                        parsed_data[key] = re.sub("smiles", "", num_match) if num_match else None
                    else:
                        parsed_data[key] = None if not match else match.group(1)
                else:
                    parsed_data[key] = match.group(1) if match else None
            return parsed_data

        parsed_data = df[column_name].apply(parse_row).apply(pd.Series)
        df = pd.concat([df, parsed_data], axis=1)
        return df

    df_1 = parse_filenames(df, "filenames")
    df_1 = df_1.loc[:, ~df_1.columns.duplicated()]
    df_1.to_csv("../data/sampling/filtering_comparison_parsed.csv", index=False)
    df_1[["model_name", "num_total_molecules", "num_filtered_molecules"]].groupby("model_name").mean()
    return (df_1,)


@app.cell
def _(df_1, df_experts):
    res_df = df_1.groupby("model_name")["smiles_after_filtering"].agg(lambda x: [item for sublist in x for item in sublist])
    res_df = res_df.apply(set)
    print(res_df.apply(len))
    res_df = res_df.apply(lambda x: x - df_experts)
    print(res_df.apply(len))
    return (res_df,)


@app.cell
def _(res_df):
    res_df.to_csv("data/sampling/xlsx_input.csv", index=True)
    return


@app.cell
def _():
    # TODO: execute script to generate excel
    return


@app.cell
def _(experts_dict):
    import os
    import shutil

    from openpyxl import Workbook
    from openpyxl.drawing.image import Image
    from openpyxl.worksheet.table import Table, TableStyleInfo
    from rdkit import Chem
    from rdkit.Chem import Draw

    def generate_excel_from_dict(model_data: dict, output_file: str, output_folder: str = "temp_images"):
        """
        Generates an Excel file with SMILES structures, images, and filter pass/fail
        columns for each filter, directly from a dictionary input.

        Args:
            model_data (dict): Dictionary of models and their SMILES DataFrames,
                where DataFrames MUST have SMILES as index and filter columns
                with boolean results.
            output_file (str): Path to save the output Excel file.
            output_folder (str): Path to the folder for temporary images.
                Defaults to 'temp_images'.

        Returns:
            None
        """
        wb = Workbook()
        os.makedirs(output_folder, exist_ok=True)
        for model_name, df in model_data.items():
            ws = wb.create_sheet(title=model_name)
            headers = ["SMILES", "Structure", "Grade (0-5)"]
            filter_columns = [col for col in df.columns if col != "smiles"]
            filter_names = [f"Passed {col}" for col in filter_columns]
            headers.extend(filter_names)
            header_row = 1
            for col_num, header in enumerate(headers, start=1):
                ws.cell(row=header_row, column=col_num, value=header)
            ws.column_dimensions["A"].width = 30
            ws.column_dimensions["B"].width = 50
            ws.column_dimensions["C"].width = 20
            start_filter_col = 4
            for i in range(len(filter_names)):
                ws.column_dimensions[chr(ord("D") + i)].width = 15
            for row in range(2, len(df) + 2):
                ws.row_dimensions[row].height = 250
            for i, row in enumerate(df.iterrows(), start=2):
                smiles = df.index[i - 2]
                data = df.iloc[i - 2]
                mol = Chem.MolFromSmiles(smiles)
                ws.cell(row=i, column=1, value=smiles)
                img_path = os.path.join(output_folder, f"{model_name}_mol_{i}.png")
                if mol:
                    Draw.MolToFile(mol, img_path)
                    img = Image(img_path)
                    ws.add_image(img, f"B{i}")
                else:
                    ws.cell(row=i, column=2, value="Invalid SMILES")
                ws.cell(row=i, column=3, value="Enter 0-5")
                filter_result_col = start_filter_col
                for filter_col_name in filter_columns:
                    passed_filter = data[filter_col_name]
                    ws.cell(row=i, column=filter_result_col, value="True" if passed_filter else "False")
                    filter_result_col = filter_result_col + 1
            last_col_letter = chr(ord("A") + len(headers) - 1)
            last_row_num = len(df) + 1
            table_range = f"A1:{last_col_letter}{last_row_num}"
            style = TableStyleInfo(name="TableStyleMedium9", showFirstColumn=False, showLastColumn=False, showRowStripes=True, showColumnStripes=False)
            tab = Table(displayName=f"Table_{model_name}", ref=table_range, tableStyleInfo=style)
            ws.add_table(tab)
        if "Sheet" in wb.sheetnames:
            del wb["Sheet"]
        wb.save(output_file)
        print(f"Excel file saved as {output_file}.")
        if os.path.exists(output_folder):
            shutil.rmtree(output_folder)
            print(f"Temporary folder '{output_folder}' has been deleted.")

    generate_excel_from_dict(experts_dict, "experts_filters_new.xlsx")
    return (Chem,)


@app.cell
def _():
    import marimo as mo

    return (mo,)


if __name__ == "__main__":
    app.run()
